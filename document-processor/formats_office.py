from __future__ import annotations
import re
from pathlib import PurePosixPath
from engine import ExtractionError

def extract_docx(engine, path):
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    from defusedxml import ElementTree as ET
    with engine.office_archive(path) as archive:
        document = Document(path)
        properties = document.core_properties
        engine.result["metadata"].update({"document_properties": {key: getattr(properties, key, None) for key in ("title","subject","author","keywords","comments","created","modified","last_modified_by","revision")}, "pagination": "DOCX has no reliable page boundaries without a layout renderer."})
        blocks = list(document.iter_inner_content())
        total = len(blocks) or 1
        def table_content(table, location):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            engine.add("table", "\n".join("\t".join(row) for row in rows), {"rows": rows}, location)
            for row_index, row in enumerate(table.rows):
                for col_index, cell in enumerate(row.cells):
                    for nested_index, nested in enumerate(cell.tables):
                        table_content(nested, {**location,"nested_cell":[row_index,col_index],"nested_table":nested_index})
        for index, block in enumerate(blocks):
            source = {"part":"word/document.xml","block":index + 1}
            engine.progress(f"Processing Word block {index + 1} of {len(blocks)}", index, len(blocks))
            if isinstance(block, Paragraph):
                style = block.style.name if block.style else ""
                kind = "heading" if style.lower().startswith(("heading","title","subtitle")) else "paragraph"
                engine.add(kind, block.text, {"style": style}, source)
            elif isinstance(block, Table): table_content(block, source)
            engine.result["units"].append({**source,"status":"processed"})
            engine.progress(f"Processed Word block {index + 1}", index + 1, len(blocks))
        for part in archive.namelist():
            if part.startswith("word/media/"):
                data = archive.read(part)
                asset = engine.asset(data, PurePosixPath(part).suffix.lstrip("."), "image", {"part": part}, {"package_part": part})
                engine.ocr(data, {"part":part,"asset_path":asset["path"]})
            elif part.startswith("word/embeddings/"):
                engine.asset(archive.read(part), "bin", "embedded_evidence", {"part":part}, {"package_part":part, "preview":False})
                engine.warning("embedded_object", "An embedded Office object is retained as evidence but is not executed or interpreted.", {"part":part})
            elif part.endswith(".rels"):
                xml = ET.fromstring(archive.read(part))
                for rel in xml:
                    if rel.attrib.get("Type","").endswith("/hyperlink"):
                        url = rel.attrib.get("Target","")
                        engine.add("url",url,{"url":url,"relationship":rel.attrib}, {"part":part}, enrich=False)
            elif part.startswith("word/") and part.endswith(".xml"):
                data = archive.read(part)
                engine.asset(data,"xml","source_xml",{"part":part},{"package_part":part,"preview":False})
                if re.search(r"(?:header|footer|footnotes|endnotes|comments)\d*\.xml$",part):
                    xml = ET.fromstring(data)
                    texts = [node.text for node in xml.iter() if node.tag.endswith("}t") and node.text is not None]
                    engine.add("paragraph","\n".join(texts),{"part_type":"supplementary"}, {"part":part})
                if part=="word/document.xml" and (b"<w:ins" in data or b"<w:del" in data):
                    engine.warning("tracked_revisions","Tracked revisions are retained in source XML; review their intended accepted/deleted state.",{"part":part})
        engine.warning("word_pagination","Word paragraph/table order is preserved. Accurate rendered page numbers require an Office layout engine.")
        if not engine.result["units"]: engine.result["units"].append({"part":"word/document.xml","status":"empty"})

def extract_xlsx(engine, path):
    from openpyxl import load_workbook
    from defusedxml import ElementTree as ET
    with engine.office_archive(path) as archive:
        workbook = load_workbook(path, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        try:
            engine.result["metadata"].update({"worksheet_count":len(workbook.worksheets),"sheet_names":workbook.sheetnames,"calculation":"Formulas are preserved but never recalculated."})
            total_cells = 0
            for number, sheet in enumerate(workbook.worksheets):
                source={"worksheet":sheet.title}
                engine.progress(f"Processing worksheet {number + 1} of {len(workbook.worksheets)}",number,len(workbook.worksheets))
                total_cells += sheet.max_row * sheet.max_column
                if total_cells > int(engine.options.get("max_cells",250000)): raise ExtractionError("Workbook exceeds the configured cell limit; increase it for a trusted workbook.")
                rows=[]; formula_missing=0; cache=cached[sheet.title]; populated=0
                for row in sheet.iter_rows():
                    values=[]
                    for cell in row:
                        value=cell.value
                        values.append(value)
                        if value is not None or cell.hyperlink:
                            populated+=1
                            record={"coordinate":cell.coordinate,"value":value,"data_type":cell.data_type,"number_format":cell.number_format}
                            if cell.data_type=="f":
                                record["formula"]=value;record["cached_value"]=cache[cell.coordinate].value
                                if record["cached_value"] is None: formula_missing+=1
                            engine.add("cell", "" if value is None else str(value),record,{**source,"cell":cell.coordinate})
                            if cell.hyperlink:
                                target=cell.hyperlink.target or cell.hyperlink.location or ""
                                engine.add("url",target,{"url":target}, {**source,"cell":cell.coordinate},enrich=False)
                    rows.append(values)
                if populated:
                    text="\n".join("\t".join("" if v is None else str(v) for v in row) for row in rows)
                    engine.add("worksheet_text",text,{"range":sheet.calculate_dimension()},source)
                    engine.add("table",text,{"rows":rows,"headers":rows[0] if rows else [],"merged_ranges":[str(x) for x in sheet.merged_cells.ranges],"named_tables":[{"name":t.name,"ref":t.ref} for t in sheet.tables.values()]},source)
                engine.result["units"].append({**source,"status":"processed","rows":sheet.max_row,"columns":sheet.max_column,"populated_cells":populated})
                if formula_missing: engine.warning("formula_cache_missing",f"{formula_missing} formula cells lack cached results; formulas and original workbook are preserved.",source)
                engine.progress(f"Processed worksheet {sheet.title}",number + 1,len(workbook.worksheets))
            for part in archive.namelist():
                if part.startswith("xl/media/"):
                    data=archive.read(part)
                    asset=engine.asset(data,PurePosixPath(part).suffix.lstrip("."),"image",{"part":part},{"package_part":part})
                    engine.ocr(data,{"part":part,"asset_path":asset["path"]})
                elif part.startswith("xl/charts/") and part.endswith(".xml"):
                    data=archive.read(part);xml=ET.fromstring(data)
                    engine.asset(data,"xml","chart_evidence",{"part":part},{"package_part":part,"preview":False})
                    text="\n".join(node.text for node in xml.iter() if node.text is not None)
                    refs=[node.text for node in xml.iter() if node.tag.endswith("}f") and node.text]
                    engine.add("chart",text,{"references":refs,"xml":data.decode("utf-8"),"cached_values":[node.text for node in xml.iter() if node.tag.endswith("}v")],"review_required":True},{"part":part},"ambiguous")
                    engine.warning("chart_not_rendered","Chart XML, series references and caches are preserved; no chart is recalculated or rendered.",{"part":part})
                elif part.startswith("xl/") and (part.endswith(".xml") or part.endswith(".rels")):
                    engine.asset(archive.read(part),"xml","source_xml",{"part":part},{"package_part":part,"preview":False})
        finally:
            workbook.close();cached.close()
