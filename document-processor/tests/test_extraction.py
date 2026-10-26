import io,json,tempfile,unittest,zipfile
from pathlib import Path
from unittest.mock import patch
from engine import Extractor,ExtractionError,doctor
import pymupdf
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from openpyxl import Workbook
from openpyxl.chart import BarChart,Reference

class ExtractionTests(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def extract(self,path,**options):return Extractor(self.root/(path.stem+'-out'),{'ocr_engine':'none',**options}).extract(path)
 def picture(self):
  im=Image.new('RGB',(1000,220),'white');ImageDraw.Draw(im).text((30,40),'Event count 97 Clicks 37',font=ImageFont.load_default(size=48),fill='black')
  b=io.BytesIO();im.save(b,format='PNG');return b.getvalue()
 def test_pdf_all_pages_text_ruled_tables_images_urls_and_source(self):
  path=self.root/'multi.pdf';d=pymupdf.open()
  p=d.new_page();p.insert_text((50,40),'Keyword ranking report',fontsize=20)
  for x in (50,250,450):p.draw_line((x,100),(x,190))
  for y in (100,130,160,190):p.draw_line((50,y),(450,y))
  for y,values in [(120,['Keyword','Rank']),(150,['Pharmacy Fairfax','3']),(180,['Pharmacy Fairfax','3'])]:
   for x,v in zip((60,260),values):p.insert_text((x,y),v)
  p.insert_link({'kind':pymupdf.LINK_URI,'from':pymupdf.Rect(50,30,200,50),'uri':'https://example.test/source'})
  p=d.new_page();p.insert_text((50,50),'Users 4,515 Clicks 2,950 CTR 0.68%');p.insert_image(pymupdf.Rect(40,100,540,210),stream=self.picture())
  p=d.new_page();p.insert_text((50,50),'Third page retained');d.save(path);d.close()
  events=[];e=Extractor(self.root/'pdf-out',{'ocr_engine':'none'},events.append);r=e.extract(path)
  self.assertEqual(3,r['metadata']['page_count']);self.assertEqual(3,len(r['units']))
  self.assertIn('Third page retained',r['full_text']);self.assertIn('4,515',r['full_text'])
  tables=[c for c in r['content'] if c['type']=='table'];self.assertTrue(tables);self.assertEqual('Rank',tables[0]['data']['rows'][0][1])
  self.assertEqual(2,sum('Pharmacy Fairfax' in str(row) for row in tables[0]['data']['rows']))
  self.assertTrue(any(c['type']=='url' and c['source']['page']==1 for c in r['content']))
  self.assertTrue(any(a['type']=='image' and a['source']['page']==2 for a in r['assets']))
  self.assertEqual(3,sum(a['type']=='page_preview' for a in r['assets']))
  self.assertTrue(all((self.root/'pdf-out'/a['path']).exists() for a in r['assets']))
  self.assertEqual(3,events[-1]['completed'])
 def test_docx_paragraphs_headings_tables_hyperlinks_and_images(self):
  path=self.root/'word.docx';d=Document();d.add_heading('Off-page activity',1);d.add_paragraph('Article submission https://example.test/article')
  t=d.add_table(rows=2,cols=2);t.cell(0,0).text='Keyword';t.cell(0,1).text='Rank';t.cell(1,0).text='Specialty Pharmacy';t.cell(1,1).text='2'
  d.add_picture(io.BytesIO(self.picture()));d.save(path);r=self.extract(path)
  self.assertIn('Specialty Pharmacy',json.dumps(r,default=str));self.assertTrue(any(c['type']=='heading' for c in r['content']))
  self.assertTrue(any(c['type']=='url' for c in r['content']));self.assertTrue(any(a['type']=='image' for a in r['assets']))
  self.assertTrue(any(w['code']=='word_pagination' for w in r['warnings']))
 def test_xlsx_all_worksheets_formulas_dates_chart_metadata(self):
  from datetime import datetime
  path=self.root/'sheets.xlsx';w=Workbook();s=w.active;s.title='Traffic';s.append(['Date','Users','URL']);s.append([datetime(2026,8,1),4515,'https://example.test/landing']);s['C2'].hyperlink='https://example.test/landing';s['B3']='=SUM(B2:B2)'
  c=BarChart();c.add_data(Reference(s,min_col=2,min_row=1,max_row=3),titles_from_data=True);s.add_chart(c,'E2');w.create_sheet('Rankings').append(['Pharmacy','3']);w.save(path)
  r=self.extract(path);self.assertEqual(2,r['metadata']['worksheet_count'])
  raw=json.dumps(r,default=str);self.assertIn('=SUM(B2:B2)',raw);self.assertIn('2026-08-01',raw);self.assertIn('Rankings',raw)
  self.assertTrue(any('chart' in a['type'] for a in r['assets']));self.assertTrue(any('formula' in w['code'] for w in r['warnings']))
 def test_csv_keeps_duplicate_rows_blanks_numeric_strings_and_utf16(self):
  path=self.root/'rows.csv';path.write_bytes('keyword,rank,ctr\nPharmacy,03,0.68%\nPharmacy,03,0.68%\n\n'.encode('utf-16'))
  r=self.extract(path);t=next(c for c in r['content'] if c['type']=='table')
  self.assertEqual(2,sum(row==['Pharmacy','03','0.68%'] for row in t['data']['rows']))
  self.assertIn('0.68%',r['full_text'])
 def test_jpeg_image_preserves_dimensions_and_private_evidence_without_ocr(self):
  from PIL import Image
  import io
  path=self.root/'source.jpg';img=Image.open(io.BytesIO(self.picture())).convert('RGB');img.save(path,'JPEG');r=self.extract(path)
  self.assertEqual('partially_extracted',r['status']);self.assertEqual(1,r['statistics']['images']);self.assertTrue(any(a['type']=='image' for a in r['assets']))
 def test_scanned_image_preserved_without_ocr_and_flagged_partial(self):
  path=self.root/'scan.png';path.write_bytes(self.picture());r=self.extract(path)
  self.assertEqual('partially_extracted',r['status']);self.assertEqual(1,r['statistics']['images']);self.assertTrue(any(w['code']=='ocr_unavailable' for w in r['warnings']))
 def test_real_ocr_when_available(self):
  if not doctor()['ocr']['available']:self.skipTest('Install optional OCR requirements')
  path=self.root/'ocr.png';path.write_bytes(self.picture());r=self.extract(path,ocr_engine='auto')
  self.assertTrue(r['statistics']['ocr_used']);self.assertIn('97',r['full_text']);self.assertIn('37',r['full_text']);self.assertIn('event',r['full_text'].lower())
 def test_corrupt_pdf_missing_format_and_page_limit_fail_honestly(self):
  bad=self.root/'bad.pdf';bad.write_bytes(b'%PDF-not-readable')
  with self.assertRaises(ExtractionError):self.extract(bad)
  path=self.root/'limit.pdf';d=pymupdf.open();d.new_page();d.new_page();d.save(path);d.close()
  with self.assertRaisesRegex(ExtractionError,'page limit'):self.extract(path,max_pages=1)
  with self.assertRaisesRegex(ExtractionError,'Unsupported'):Extractor(self.root/'bad').extract(path,'exe')
 def test_office_archive_rejects_traversal_macros_and_xml_entities(self):
  for name,data in [('word/../danger','x'),('word/vbaProject.bin','x'),('word/document.xml','<!DOCTYPE doc [<!ENTITY a "x">]><doc/>')]:
   path=self.root/'unsafe.docx'
   with zipfile.ZipFile(path,'w') as z:z.writestr(name,data)
   with self.assertRaises(ExtractionError):Extractor(self.root/'unsafe').office_archive(path)
 def test_partial_content_survives_later_parser_failure(self):
  path=self.root/'partial.csv';path.write_text('a,b\n1,2',encoding='utf8')
  def failure(e,p):e.add('text','Preserved first segment');raise ValueError('failure')
  with patch('formats_data.extract_csv',failure):
   r=self.extract(path);self.assertEqual('partially_extracted',r['status']);self.assertIn('Preserved',r['full_text']);self.assertTrue(r['errors'])
 def test_legacy_xls_all_sheets_cached_values_and_limit_warning(self):
  import xlwt
  path=self.root/'legacy.xls';book=xlwt.Workbook()
  a=book.add_sheet('Users');a.write(0,0,'Users');a.write(1,0,4515)
  b=book.add_sheet('Keywords');b.write(0,0,'Specialty pharmacy');b.write(0,1,3);book.save(str(path))
  r=self.extract(path);self.assertEqual(2,r['metadata']['worksheet_count']);self.assertIn('4515',r['full_text']);self.assertIn('Specialty pharmacy',r['full_text'])
  self.assertTrue(any(w['code']=='legacy_excel_limit' for w in r['warnings']))
 def test_doctor_reports_real_capabilities(self):
  h=doctor();self.assertTrue(h['formats']['pdf']);self.assertTrue(h['formats']['csv']);self.assertIn('available',h['ocr'])


class PipelineRepairTests(unittest.TestCase):
 def test_repeated_image_ocr_is_cached_without_losing_source_occurrences(self):
  if not doctor()['ocr']['available']:self.skipTest('OCR unavailable')
  image=Image.new('RGB',(900,200),'white');ImageDraw.Draw(image).text((30,40),'Pharmacy Events 37',font=ImageFont.load_default(size=48),fill='black')
  data=io.BytesIO();image.save(data,format='PNG')
  with tempfile.TemporaryDirectory() as tmp:
   e=Extractor(tmp,{'ocr_engine':'auto'})
   e.ocr(data.getvalue(),{'page':1});e.ocr(data.getvalue(),{'page':2})
   self.assertEqual(1,e.ocr_cache_hits)
   self.assertEqual([1,2],[c['source']['page'] for c in e.result['content'] if c['type']=='ocr_text'])
   self.assertTrue(all('37' in c['text'] for c in e.result['content'] if c['type']=='ocr_text'))
 def test_table_parser_failure_does_not_skip_page_rendering_or_later_pages(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);path=root/'text.pdf';d=pymupdf.open()
   for n in range(2):d.new_page().insert_text((40,40),'Native text on page '+str(n+1))
   d.save(path);d.close()
   with patch('pymupdf.Page.find_tables',side_effect=RuntimeError('parser failure')):
    r=Extractor(root/'out',{'ocr_engine':'none'}).extract(path)
   self.assertEqual(2,sum(a['type']=='page_preview' for a in r['assets']))
   self.assertIn('page 2',r['full_text']);self.assertEqual('partially_extracted',r['status'])
   self.assertTrue(any(e['code']=='table_failed' for e in r['errors']))


class NativeCellTests(unittest.TestCase):
 def test_native_table_cells_keep_url_underscores_and_parser_evidence(self):
  source=next((p for p in (Path(__file__).resolve().parents[2]/'reference-files').glob('*.pdf') if 'Website_Analytics' in p.name),None)
  if not source:self.skipTest('Reference PDF not available')
  with tempfile.TemporaryDirectory() as tmp:
   r=Extractor(tmp,{'ocr_engine':'none'}).extract(source)
   tables=[c for c in r['content'] if c['type']=='table' and c['source']['page']==4]
   self.assertTrue(tables)
   self.assertTrue(any('Leuprolide_Dose_Card_v1.pdf' in str(row) for t in tables for row in t['data']['rows']))
   self.assertTrue(any('utm_source=google&utm_medium=organic&utm_campaign=gbp' in str(row) for t in tables for row in t['data']['rows']))
   self.assertTrue(all('parser_rows' in t['data'] and 'cell_bboxes' in t['data'] for t in tables))
   for table in tables:
    for original,row in zip(table['data']['parser_rows'],table['data']['rows']):
     for before,after in zip(original,row):
      if before is None or not str(before).strip():self.assertEqual(before,after)
