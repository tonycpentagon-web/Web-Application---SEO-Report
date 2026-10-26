"""Generate synthetic source fixtures only; never opens original client files."""
from pathlib import Path
import io
import pymupdf as fitz
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from openpyxl import Workbook

def generate(root):
 root=Path(root);root.mkdir(parents=True,exist_ok=True)
 image=Image.new('RGB',(1200,500),'white');draw=ImageDraw.Draw(image)
 font_path=next((p for p in [Path('C:/Windows/Fonts/arial.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')] if p.exists()),None)
 font=ImageFont.truetype(str(font_path),36) if font_path else ImageFont.load_default()
 draw.multiline_text((40,40),'Event name    Event count\ncontact_us    10\nphone_tap     5\nGrand total   15',fill='black',font=font,spacing=20)
 image.save(root/'key events.png');image.save(root/'screenshot.jpeg')
 bio=io.BytesIO();image.save(bio,format='PNG')
 pdf=fitz.open()
 texts=['Total users\n800\nNew users\n700\nViews\n900\nClicks\n90\nImpressions\n2000\nSite CTR\n4.5%', 'KEYWORDS\nRANKING\nSynthetic treatment\n1st Page','ARTICLE SUBMISSION\nhttps://example.test/synthetic-article','AI VISIBILITY\nSynthetic query, evidence only','Eventname\nEventcount\n1.\ncontact_us\n10\n2.\nphone_tap\n5\nGrand total\n15']
 for i,text in enumerate(texts):
  page=pdf.new_page();page.insert_text((50,60),text,fontsize=14)
  if i==3:page.insert_image(fitz.Rect(50,200,550,410),stream=bio.getvalue())
 pdf.save(root/'Keyword_Ranking.pdf');pdf.close()
 d=Document();d.add_heading('ARTICLE SUBMISSION',0);d.add_paragraph('https://example.test/second-synthetic-article');d.save(root/'activities.docx')
 wb=Workbook();ws=wb.active;ws.append(['Keyword','Ranking']);ws.append(['Synthetic other','2nd Page']);wb.save(root/'keywords.xlsx')
 (root/'channels.csv').write_text('Channel,Total users,New users\nOrganic Search,500,450\nDirect,300,250\n')
 # Exactly five fixture files for the legacy batch test; supplemental formats live separately.
 extra=root/'supplemental';extra.mkdir(exist_ok=True);(root/'screenshot.jpeg').rename(extra/'screenshot.jpeg')
 (extra/'large-keywords.csv').write_text('Keyword,Ranking\n'+''.join(f'Synthetic keyword {i},2nd Page\n' for i in range(1000)))
 scanned=fitz.open();p=scanned.new_page();p.insert_image(fitz.Rect(30,30,570,300),stream=bio.getvalue());scanned.save(extra/'scanned.pdf');scanned.close()
 return root
if __name__=='__main__':
 import sys
 generate(sys.argv[1])
